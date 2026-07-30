"""
안동시 시내버스 운행시간표 엑셀 → 정규화 JSON

각 시트는 사람이 보라고 만든 표라 컬럼 구조가 제각각이다.
완전 파싱을 시도하지 않고, 확실하게 뽑을 수 있는 것만 수확한다.

  수확 대상: 노선 × 정류장(컬럼) 별 첫차 / 막차 / 운행 횟수
  수확 제외: 경유지 메모, 조건부 운행 규칙, 차량 종류 등

"없는 건 없는 대로 두고, 있는 것만 정확히" 가 원칙.
"""
import openpyxl, datetime, json, re, os, glob

# ── 시간 정규화 ────────────────────────────────────────
TIME_RE = re.compile(r'(?<!\d)([01]?\d|2[0-3])\s*:\s*([0-5]\d)(?!\d)')

def to_minutes(v):
    """셀 값에서 HH:MM을 뽑아 분 단위 정수로. 실패하면 None."""
    if v is None:
        return None
    if isinstance(v, datetime.datetime):
        return v.hour * 60 + v.minute
    if isinstance(v, datetime.time):
        return v.hour * 60 + v.minute
    if isinstance(v, str):
        m = TIME_RE.search(v)
        if m:
            return int(m.group(1)) * 60 + int(m.group(2))
    return None

def hhmm(mins):
    return f"{mins // 60:02d}:{mins % 60:02d}"

# ── 컬럼 라벨 정리 ─────────────────────────────────────
SKIP_LABELS = {'구분', '경유', '비고', '경유지', '', '순번'}

def clean_label(v):
    if v is None:
        return ''
    return re.sub(r'\s+', '', str(v).replace('\n', ''))

NOISE_RE = re.compile(r'시행|^\d{2}[.\-]\d{1,2}|^구\d|^\d{4}$')

def is_stop_column(label):
    if label in SKIP_LABELS:
        return False
    if len(label) > 14:          # 문장이 들어간 셀
        return False
    if re.fullmatch(r'\d+', label):
        return False
    if NOISE_RE.search(label):   # "24.1.5.시행" 같은 메타 문구
        return False
    return True

# ── 시트 메타 추출 ─────────────────────────────────────
DAY_PATTERNS = [
    (re.compile(r'토\s*,?\s*일\s*,?\s*공'), 'weekend_holiday'),
    (re.compile(r'\(일\s*,?\s*공\)'),        'sunday_holiday'),
    (re.compile(r'\(토\)'),                  'saturday'),
    (re.compile(r'\(평일\)'),                'weekday'),
]

def day_type(sheet_name):
    for pat, key in DAY_PATTERNS:
        if pat.search(sheet_name):
            return key
    return 'all'      # 요일 구분 표기가 없으면 매일 동일로 간주

ROUTE_RE = re.compile(r'(급행\s*\d|순환\s*\d(?:-\d)?|\d{3}|\d{2,3}번)')

def route_code(sheet_name):
    m = ROUTE_RE.search(sheet_name.replace(' ', ''))
    return m.group(1).replace('번', '') if m else sheet_name.split('(')[0].strip()

EFF_RE = re.compile(r'(\d{2})\s*[.\-]\s*(\d{1,2})\s*[.\-]\s*(\d{1,2})')

def effective_date(rows):
    for r in rows[:6]:
        for c in r:
            if isinstance(c, str) and '시행' in c:
                m = EFF_RE.search(c)
                if m:
                    y, mo, d = m.groups()
                    return f"20{y}-{int(mo):02d}-{int(d):02d}"
    return None

def sheet_notes(rows, header_idx):
    """헤더 위쪽의 안내 문구(조건부 운행 등)를 모은다."""
    notes = []
    for r in rows[:header_idx]:
        for c in r:
            if isinstance(c, str):
                t = c.strip().replace('\n', ' ')
                if len(t) > 6 and '시행' not in t:
                    notes.append(t)
    return notes

# ── 시트 1개 파싱 ──────────────────────────────────────
def parse_sheet(ws):
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return None

    header_idx = None
    for i, r in enumerate(rows[:10]):
        if any(isinstance(c, str) and c.strip() == '구분' for c in r):
            header_idx = i
            break

    # '구분' 열이 없는 시트(113·114·513·111탄력 등) 대응:
    # 시각이 2개 이상 등장하기 시작하는 행의 바로 위를 헤더로 본다.
    if header_idx is None:
        for i, r in enumerate(rows[:12]):
            if sum(1 for c in r if to_minutes(c) is not None) >= 2:
                header_idx = max(i - 1, 0)
                break
    if header_idx is None:
        header_idx = 0

    header = [clean_label(c) for c in rows[header_idx]]

    # 헤더가 2행에 걸친 경우(급행2 등) 아랫줄 라벨로 보완
    if header_idx + 1 < len(rows):
        sub = [clean_label(c) for c in rows[header_idx + 1]]
        header = [h if h else sub[i] if i < len(sub) else ''
                  for i, h in enumerate(header)]

    stops = []
    for col, label in enumerate(header):
        if not is_stop_column(label):
            continue
        mins = []
        for r in rows[header_idx + 1:]:
            if col < len(r):
                m = to_minutes(r[col])
                if m is not None:
                    mins.append(m)
        if len(mins) >= 2:
            mins.sort()
            stops.append({
                'stop': label,
                'first': hhmm(mins[0]),
                'last': hhmm(mins[-1]),
                'runs': len(mins),
                'times': [hhmm(m) for m in mins],
            })

    if not stops:
        return None

    return {
        'effectiveFrom': effective_date(rows),
        'notes': sheet_notes(rows, header_idx),
        'stops': stops,
    }

# ── 전체 처리 ─────────────────────────────────────────
def build(files):
    out = []
    for path in files:
        wb = openpyxl.load_workbook(path, data_only=True)
        src = os.path.basename(path)
        for name in wb.sheetnames:
            parsed = parse_sheet(wb[name])
            if not parsed:
                continue
            out.append({
                'sheet': name,
                'routeCode': route_code(name),
                'dayType': day_type(name),
                'source': src,
                **parsed,
            })
    return out


if __name__ == '__main__':
    HERE = os.path.dirname(os.path.abspath(__file__))
    SRC = os.path.join(HERE, 'source')
    OUT = os.path.join(HERE, 'output', 'timetable-raw.json')

    files = sorted(glob.glob(os.path.join(SRC, '*.xlsx')))
    if not files:
        raise SystemExit(
            f'엑셀 파일이 없습니다: {SRC}\n'
            '안동시 공공데이터포털에서 운행시간표를 내려받아 scripts/source/ 에 넣으세요.'
        )

    data = build(files)
    print(f'파싱된 시트: {len(data)}개')
    print(f"수확된 정류장 컬럼: {sum(len(d['stops']) for d in data)}개")

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    with open(OUT, 'w', encoding='utf-8') as f:
        json.dump(data, f, ensure_ascii=False, indent=2)
    print(f'저장: {OUT}')
